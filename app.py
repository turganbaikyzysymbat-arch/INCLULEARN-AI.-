import io
import json
import os
import re
import sqlite3
import urllib.request
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request


BASE_DIR = Path(__file__).resolve().parent
DB_DIR = BASE_DIR / "instance"
DB_PATH = DB_DIR / "learn4all.db"
DB_DIR.mkdir(exist_ok=True)

app = Flask(__name__, template_folder="templates", static_folder="static")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

STOP_WORDS = set(
    "және мен бұл үшін туралы немесе бір оның болып арқылы да де әрі қалай деген "
    "the and of to a in is it that for with as on are was this be by an or from at "
    "и в на это как для что из с по".split()
)

LANGUAGE_NAMES = {
    "kk": "қазақ тілі",
    "ru": "русский язык",
    "en": "English",
}


@contextmanager
def db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    finally:
        connection.close()


def init_db():
    DB_DIR.mkdir(exist_ok=True)
    with db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS materials (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                source_text TEXT NOT NULL,
                file_name TEXT,
                language TEXT DEFAULT 'kk',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS adaptations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL UNIQUE,
                result_json TEXT NOT NULL,
                engine TEXT NOT NULL DEFAULT 'local',
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL,
                question TEXT NOT NULL,
                answer TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS progress (
                material_id INTEGER PRIMARY KEY,
                completed INTEGER NOT NULL DEFAULT 0,
                score INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            """
        )


def now():
    return datetime.now().isoformat(timespec="seconds")


def clean_text(value):
    value = str(value or "").replace("\x00", " ").replace("\r", "")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def split_sentences(text):
    return [
        item.strip()
        for item in re.split(
            r"(?<=[.!?])\s+(?=[А-ЯӘҒҚҢӨҰҮҺІA-Z0-9])", clean_text(text)
        )
        if item.strip()
    ]


def get_words(text):
    return re.findall(
        r"[а-яәіңғқөұүһіa-z0-9'-]+", clean_text(text).lower(), flags=re.I
    )


def keywords(text, limit=8):
    counts = {}
    for word in get_words(text):
        if len(word) >= 4 and word not in STOP_WORDS:
            counts[word] = counts.get(word, 0) + 1
    return [
        word
        for word, _ in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[
            :limit
        ]
    ]


def select_summary(text, limit=4):
    sentence_list = split_sentences(text)
    if len(sentence_list) <= limit:
        return sentence_list
    terms = keywords(text, limit)
    scored = []
    for position, sentence in enumerate(sentence_list):
        score = sum(2 for term in terms if term in sentence.lower())
        score += 1 if position == 0 else 0
        scored.append((score, position, sentence))
    return [
        row[2]
        for row in sorted(sorted(scored, reverse=True)[:limit], key=lambda row: row[1])
    ]


def local_adaptation(text, language="kk"):
    source = clean_text(text)
    terms = keywords(source)
    summary_sentences = select_summary(source)
    topic = ", ".join(terms[:3]) or "негізгі тақырып"
    short = "\n".join(
        f"• {sentence[:180]}{'…' if len(sentence) > 180 else ''}"
        for sentence in summary_sentences
    )
    steps = "\n".join(
        f"{index}. {sentence}"
        for index, sentence in enumerate(split_sentences(source)[:8], 1)
    ) or "1. Мәтінді оқыңыз.\n2. Негізгі ұғымдарды белгілеңіз.\n3. Мысал арқылы тексеріңіз."
    examples = "\n".join(
        f"{index}. «{term}» ұғымын күнделікті өмірдегі жағдаймен байланыстырып, өз мысалыңызды жазыңыз."
        for index, term in enumerate(terms[:4], 1)
    ) or "1. Тақырыпқа қатысты қарапайым жағдай ойластырыңыз.\n2. Оны негізгі ұғыммен байланыстырыңыз."
    tasks = (
        f"1. «{terms[0] if terms else 'негізгі ұғым'}» ұғымын өз сөзіңізбен түсіндіріңіз.\n"
        f"2. {terms[1] if len(terms) > 1 else 'тақырып'} бойынша бір мысал келтіріңіз.\n"
        "3. Мәтіндегі екі негізгі ойдың байланысын жазыңыз.\n"
        "4. Жауабыңыздан негізгі терминдерді белгілеңіз."
    )
    plain = re.sub(r"\([^)]*\)", "", source).replace(";", ".")
    plain = "\n\n".join(
        sentence[:200] + ("…" if len(sentence) > 200 else "")
        for sentence in split_sentences(plain)
    )
    nodes = (terms or ["Негізгі ұғым", "Ереже", "Мысал", "Нәтиже"])[:5]
    return {
        "simplified": short or source[:600],
        "steps": steps,
        "examples": examples,
        "diagram": {
            "title": topic,
            "nodes": nodes,
            "edges": [[nodes[index], nodes[index + 1]] for index in range(len(nodes) - 1)],
        },
        "audio": f"{topic}. " + " ".join(summary_sentences),
        "large": plain or source,
        "summary": " ".join(summary_sentences) or source[:500],
        "tasks": tasks,
        "keywords": terms,
        "explanation": (
            f"Бұл материалдың негізгі бағыты: {topic}. Алдымен анықтаманы түсініңіз, "
            "кейін оның қолданылуын мысалмен тексеріңіз."
        ),
        "language": language,
    }


def ai_configured():
    return bool(os.getenv("OPENAI_API_KEY") or os.getenv("AI_INTEGRATIONS_OPENAI_API_KEY"))


def ask_model(system, prompt):
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("AI_INTEGRATIONS_OPENAI_API_KEY")
    if not api_key:
        return None
    endpoint = (
        os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
        + "/chat/completions"
    )
    payload = json.dumps(
        {
            "model": os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
            "temperature": 0.2,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
        }
    ).encode()
    model_request = urllib.request.Request(
        endpoint,
        data=payload,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
        method="POST",
    )
    with urllib.request.urlopen(model_request, timeout=45) as response:
        data = json.loads(response.read().decode())
    return data["choices"][0]["message"]["content"]


def extract_file(file):
    extension = Path(file.filename or "").suffix.lower()
    allowed = {".txt", ".md", ".csv", ".html", ".pdf", ".docx", ".pptx"}
    if extension not in allowed:
        raise ValueError("PDF, DOCX, PPTX, TXT, MD, CSV немесе HTML файлдарын қолданыңыз.")
    data = file.read()
    if extension in {".txt", ".md", ".csv", ".html"}:
        return clean_text(data.decode("utf-8", errors="ignore"))
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return clean_text("\n\n".join(page.extract_text() or "" for page in reader.pages))
    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        return clean_text("\n\n".join(paragraph.text for paragraph in document.paragraphs))
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides = []
    for slide in presentation.slides:
        slide_text = [
            shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        ]
        if slide_text:
            slides.append(" ".join(slide_text))
    return clean_text("\n\n".join(slides))


def local_answer(question, source):
    query = keywords(question, 8)
    matched = []
    for position, sentence in enumerate(split_sentences(source)):
        score = sum(1 for term in query if term in sentence.lower())
        if score:
            matched.append((score, -position, sentence))
    matched = [row[2] for row in sorted(matched, reverse=True)[:3]]
    if not matched:
        return "Бұл сұраққа жүктелген материалдан нақты жауап табылмады. Сұрақты мәтіндегі негізгі терминдермен нақтылап көріңіз."
    return "Материалға сүйенген жауап:\n\n" + "\n".join(f"• {sentence}" for sentence in matched)


def material_json(row, adaptation=None, progress=None, preview=False):
    result = dict(row)
    if preview:
        result["source_text"] = result["source_text"][:180]
    result["adaptation"] = json.loads(adaptation["result_json"]) if adaptation else None
    result["progress"] = dict(progress) if progress else {
        "completed": 0,
        "score": 0,
    }
    return result


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "database": DB_PATH.exists(), "ai_configured": ai_configured()})


@app.get("/api/materials")
def list_materials():
    query = clean_text(request.args.get("q", ""))
    with db() as connection:
        if query:
            rows = connection.execute(
                "SELECT * FROM materials WHERE title LIKE ? OR source_text LIKE ? "
                "ORDER BY updated_at DESC LIMIT 100",
                (f"%{query}%", f"%{query}%"),
            ).fetchall()
        else:
            rows = connection.execute(
                "SELECT * FROM materials ORDER BY updated_at DESC LIMIT 100"
            ).fetchall()
        materials = []
        for row in rows:
            progress = connection.execute(
                "SELECT completed, score, updated_at FROM progress WHERE material_id = ?",
                (row["id"],),
            ).fetchone()
            materials.append(material_json(row, progress=progress, preview=True))
    return jsonify({"materials": materials})


@app.get("/api/materials/<int:material_id>")
def get_material(material_id):
    with db() as connection:
        row = connection.execute(
            "SELECT * FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        adaptation = connection.execute(
            "SELECT * FROM adaptations WHERE material_id = ?", (material_id,)
        ).fetchone()
        progress = connection.execute(
            "SELECT completed, score, updated_at FROM progress WHERE material_id = ?",
            (material_id,),
        ).fetchone()
    if not row:
        return jsonify({"error": "Материал табылмады."}), 404
    return jsonify(material_json(row, adaptation, progress))


@app.post("/api/materials")
def create_material():
    try:
        body = request.get_json(silent=True) if request.is_json else {}
        title = clean_text(request.form.get("title") or body.get("title", ""))
        language = request.form.get("language") or body.get("language") or "kk"
        file = request.files.get("file")
        text = request.form.get("text", "") or body.get("text", "")
        file_name = file.filename if file else ""
        if file:
            text = extract_file(file)
        text = clean_text(text)
        if not text:
            return jsonify({"error": "Материал мәтінін енгізіңіз немесе файл жүктеңіз."}), 400
        if not title:
            title = Path(file_name).stem if file_name else "Жаңа оқу материалы"
        timestamp = now()
        with db() as connection:
            cursor = connection.execute(
                "INSERT INTO materials(title, source_text, file_name, language, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (title, text, file_name, language, timestamp, timestamp),
            )
            material_id = cursor.lastrowid
        result = await_adaptation(material_id, text, language)
        return jsonify(
            {"id": material_id, "title": title, "source_text": text, "adaptation": result}
        ), 201
    except Exception as error:
        return jsonify({"error": str(error)}), 400


def await_adaptation(material_id, text, language):
    result = local_adaptation(text, language)
    engine = "local"
    if ai_configured():
        try:
            prompt = (
                "Берілген материалды инклюзивті оқытуға бейімде. Тек материалдағы ақпаратты пайдалан. "
                f"Нәтижені {LANGUAGE_NAMES.get(language, 'қазақ тілі')} тілінде бер. "
                "Төмендегі JSON кілттерін сақта: simplified, steps, examples, explanation, summary, tasks, audio."
                f"\nМатериал:\n{text}"
            )
            response = ask_model(
                "Сен IncluLearn AI платформасындағы Ayla AI ассистентісің. Қысқа, түсінікті жауап бер.",
                prompt,
            )
            parsed = json.loads(response)
            result.update(
                {key: parsed[key] for key in parsed if key in result and isinstance(parsed[key], str)}
            )
            engine = "ai"
        except Exception:
            engine = "local"
    with db() as connection:
        connection.execute(
            "INSERT INTO adaptations(material_id, result_json, engine, created_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(material_id) DO UPDATE SET result_json=excluded.result_json, "
            "engine=excluded.engine, created_at=excluded.created_at",
            (material_id, json.dumps(result, ensure_ascii=False), engine, now()),
        )
        connection.execute(
            "UPDATE materials SET updated_at = ? WHERE id = ?", (now(), material_id)
        )
    return result


@app.put("/api/materials/<int:material_id>")
def update_material(material_id):
    body = request.get_json(silent=True) or {}
    title = clean_text(body.get("title"))
    text = clean_text(body.get("text"))
    if not title or not text:
        return jsonify({"error": "Атауы мен мәтіні бос болмауы керек."}), 400
    with db() as connection:
        exists = connection.execute(
            "SELECT id FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        if not exists:
            return jsonify({"error": "Материал табылмады."}), 404
        connection.execute(
            "UPDATE materials SET title = ?, source_text = ?, updated_at = ? WHERE id = ?",
            (title, text, now(), material_id),
        )
    result = await_adaptation(material_id, text, body.get("language", "kk"))
    return jsonify({"id": material_id, "title": title, "source_text": text, "adaptation": result})


@app.delete("/api/materials/<int:material_id>")
def delete_material(material_id):
    with db() as connection:
        cursor = connection.execute("DELETE FROM materials WHERE id = ?", (material_id,))
    if cursor.rowcount == 0:
        return jsonify({"error": "Материал табылмады."}), 404
    return jsonify({"ok": True})


@app.post("/api/materials/<int:material_id>/progress")
def save_progress(material_id):
    body = request.get_json(silent=True) or {}
    completed = 1 if body.get("completed") else 0
    score = max(0, min(100, int(body.get("score", 0))))
    with db() as connection:
        exists = connection.execute(
            "SELECT id FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        if not exists:
            return jsonify({"error": "Материал табылмады."}), 404
        connection.execute(
            "INSERT INTO progress(material_id, completed, score, updated_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(material_id) DO UPDATE SET completed=excluded.completed, "
            "score=excluded.score, updated_at=excluded.updated_at",
            (material_id, completed, score, now()),
        )
    return jsonify({"completed": completed, "score": score})


@app.post("/api/materials/<int:material_id>/ask")
def ask_about_material(material_id):
    body = request.get_json(silent=True) or {}
    question = clean_text(body.get("question"))
    if not question:
        return jsonify({"error": "Сұрақ жазыңыз."}), 400
    with db() as connection:
        material = connection.execute(
            "SELECT * FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
    if not material:
        return jsonify({"error": "Материал табылмады."}), 404
    answer = local_answer(question, material["source_text"])
    engine = "local"
    if ai_configured():
        try:
            answer = ask_model(
                "Сен Ayla AI, IncluLearn AI платформасындағы ақылды оқу ассистентісің. Оқушыға мұғалім сияқты түсінікті жауап бер. "
                "Тек берілген материалға сүйен. Жауапты оқушыға түсінікті, толық және құрылымды бер. "
                "Алдымен қысқаша негізгі жауап бер, кейін түсіндіру, мысал және маңызды терминдерді көрсет. "
                "Егер ақпарат материалда болмаса, ойдан қоспай, оны анық айт. "
                f"Жауапты {LANGUAGE_NAMES.get(material['language'], 'қазақ тілі')} тілінде бер.",
                f"Материал:\n{material['source_text']}\n\nСұрақ:\n{question}",
            )
            engine = "ai"
        except Exception:
            pass
    with db() as connection:
        connection.execute(
            "INSERT INTO questions(material_id, question, answer, created_at) VALUES (?, ?, ?, ?)",
            (material_id, question, answer, now()),
        )
    return jsonify({"answer": answer, "engine": engine})


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "3000")), debug=False)
